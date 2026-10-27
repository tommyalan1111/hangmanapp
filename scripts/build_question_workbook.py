# -*- coding: utf-8 -*-
import json, math, re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"E:\Work\Hangman app")
data = json.loads((ROOT / ".tmp-puzzles.json").read_text(encoding="utf-8"))

SHEETS = [
    ("Movies", "movies", "Movie Name", "C9A227"),
    ("TV Series", "tv", "Show Name", "2F6F7E"),
    ("Actors", "actors", "Actor Name", "B5651D"),
    ("Actresses", "actresses", "Actress Name", "B04A6E"),
    ("Characters", "characters", "Character Name", "6E5AA8"),
    ("Directors", "directors", "Director Name", "3F7D5A"),
    ("Film Families", "families", "Family Name", "A06A2C"),
]

DIFFICULTY = {1: "Easy", 2: "Medium", 3: "Hard"}
DIFF_FILL = {
    "Easy": PatternFill("solid", fgColor="E7F4EA"),
    "Medium": PatternFill("solid", fgColor="FFF4D6"),
    "Hard": PatternFill("solid", fgColor="FBE4E1"),
}
DIFF_FONT = {
    "Easy": "2E6B3D",
    "Medium": "8A5A00",
    "Hard": "A63B32",
}

SMALL = {"a", "an", "and", "as", "at", "but", "by", "for", "from", "in", "into", "nor", "of", "on", "or", "the", "to", "with", "vs", "via"}

def title_name(value: str) -> str:
    words = value.replace("’", "'").split(" ")
    out = []
    for i, word in enumerate(words):
        if not word:
            out.append(word)
            continue
        titled = word.title()
        titled = re.sub(r"'S\b", "'s", titled)
        low = titled.lower()
        prev = out[-1] if out else ""
        after_colon = prev.endswith(":") or prev.endswith("—") or prev.endswith("–")
        if 0 < i < len(words) - 1 and low in SMALL and not after_colon:
            titled = low
        out.append(titled)
    return " ".join(out)

HEADER_FILL = PatternFill("solid", fgColor="1C202B")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="F4C95D")
BODY_FONT = Font(name="Calibri", size=10.5, color="22252D")
ID_FONT = Font(name="Consolas", size=9.5, color="555B69")
NAME_FONT = Font(name="Calibri", size=10.5, bold=True, color="1B1F28")
ANSWER_FONT = Font(name="Consolas", size=10, bold=True, color="8A6415")
THIN = Side(style="thin", color="E1E3E8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
wb.remove(wb.active)
wb.properties.title = "Reel Hangman Question Bank"
wb.properties.subject = "Movies, TV, actors, actresses, characters, directors, and film families"
wb.properties.creator = "Reel Hangman"
wb.properties.description = "Question and answer bank exported from the Reel Hangman app data."

for sheet_name, category, name_header, tab_color in SHEETS:
    rows = [p for p in data if p["category"] == category]
    rows.sort(key=lambda p: (p["difficulty"], p["answer"].casefold()))
    ws = wb.create_sheet(sheet_name)
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = "A2"

    headers = ["ID", name_header, "Question", "Answer", "Difficulty"]
    ws.append(headers)
    ws.row_dimensions[1].height = 28
    for col, cell in enumerate(ws[1], 1):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center" if col == 5 else "left", vertical="center")
        cell.border = BORDER

    for idx, puzzle in enumerate(rows, start=2):
        difficulty = DIFFICULTY[puzzle["difficulty"]]
        name = title_name(puzzle["answer"])
        ws.cell(row=idx, column=1, value=puzzle["id"])
        ws.cell(row=idx, column=2, value=name)
        ws.cell(row=idx, column=3, value=puzzle["clue"])
        ws.cell(row=idx, column=4, value=puzzle["answer"])
        ws.cell(row=idx, column=5, value=difficulty)

        for col in range(1, 6):
            cell = ws.cell(row=idx, column=col)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=col in (2, 3, 4), horizontal="center" if col == 5 else "left")
            if col == 1:
                cell.font = ID_FONT
            elif col == 2:
                cell.font = NAME_FONT
            elif col == 4:
                cell.font = ANSWER_FONT
            elif col == 5:
                cell.font = Font(name="Calibri", size=10, bold=True, color=DIFF_FONT[difficulty])
                cell.fill = DIFF_FILL[difficulty]
            else:
                cell.font = BODY_FONT

        est_lines = max(
            math.ceil(len(puzzle["clue"]) / 68),
            math.ceil(len(name) / 30),
            math.ceil(len(puzzle["answer"]) / 28),
            1,
        )
        ws.row_dimensions[idx].height = max(26, 17 + (est_lines - 1) * 14)

    widths = {"A": 38, "B": 32, "C": 72, "D": 31, "E": 12}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    last = len(rows) + 1
    ws.auto_filter.ref = f"A1:E{last}"
    ws.print_area = f"A1:E{last}"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = f"{sheet_name}  ·  &P of &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "777777"

out = ROOT / "docs" / "Reel-Hangman-Question-Bank.xlsx"
wb.save(out)
print(out)
print(f"sheets={len(wb.sheetnames)} rows={sum(len([p for p in data if p['category']==c]) for _,c,_,_ in SHEETS)}")
for sheet_name, category, _, _ in SHEETS:
    print(f"  {sheet_name}: {sum(1 for p in data if p['category']==category)}")
